#!/usr/bin/env python3
"""Export any UTF-8 CSV table to XLSX while preserving the exact header schema."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path


def export_csv_to_xlsx(input_csv: Path, output_xlsx: Path, sheet_name: str = "table") -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("XLSX export requires openpyxl") from exc

    with input_csv.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        rows = list(reader)
    if not rows or not rows[0]:
        raise ValueError("CSV must contain a header row")
    width = len(rows[0])
    if any(len(row) != width for row in rows[1:]):
        raise ValueError("CSV contains rows with inconsistent column counts")

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31] or "table"
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(rows[0], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 12), 30)
    workbook.save(output_xlsx)
    return len(rows) - 1


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sheet", default="table")
    args = parser.parse_args()
    count = export_csv_to_xlsx(args.input, args.output, args.sheet)
    print(f"wrote {count} data rows to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
