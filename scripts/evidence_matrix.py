#!/usr/bin/env python3
"""Create, validate, and export the canonical evidence-matrix schema.

This utility validates provenance and traceability metadata. It does not judge
whether a scientific claim is true; full-text inspection remains a human/agent
evidence task.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Iterable

MATRIX_COLUMNS = [
    "evidence_id",
    "citation_key",
    "source_title",
    "source_type",
    "provenance_status",
    "full_text_checked",
    "locator",
    "study_design",
    "model_system",
    "intervention",
    "intervention_formulation",
    "comparator",
    "endpoint",
    "timepoint",
    "parameter_name",
    "parameter_value",
    "unit",
    "replicate_definition",
    "statistical_method",
    "quality_domains",
    "directness",
    "consistency",
    "applicability",
    "confidence",
    "parameter_status",
    "design_implication",
    "limitations",
    "review_decision",
    "reviewer",
    "reviewed_at",
]

PROVENANCE = {
    "FULL-TEXT VERIFIED",
    "SUPPLEMENT VERIFIED",
    "ABSTRACT ONLY",
    "SECONDARY CITATION",
    "DATABASE VERIFIED",
    "GUIDELINE VERIFIED",
    "UNVERIFIED",
}
PARAMETER_STATUS = {
    "EVIDENCE-SUPPORTED",
    "INDIRECT-EVIDENCE",
    "PILOT-REQUIRED",
    "ASSUMPTION",
    "USER-DECISION",
    "UNRESOLVED",
    "NOT-APPLICABLE",
}
CONFIDENCE = {"HIGH", "MODERATE", "LOW", "INSUFFICIENT"}
REVIEW_DECISIONS = {"INCLUDE", "CONTEXT-ONLY", "PILOT-ONLY", "EXCLUDE", "UNRESOLVED"}


def _truthy(value: str) -> bool:
    return value.strip().lower() in {"true", "yes", "y", "1"}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError("CSV has no header row")
        missing = [column for column in MATRIX_COLUMNS if column not in reader.fieldnames]
        if missing:
            raise ValueError(f"CSV is missing required columns: {', '.join(missing)}")
        return [{column: (row.get(column) or "").strip() for column in MATRIX_COLUMNS} for row in reader]


def write_csv(path: Path, rows: Iterable[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=MATRIX_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in MATRIX_COLUMNS})


def write_xlsx(path: Path, rows: Iterable[dict[str, str]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("XLSX export requires openpyxl") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "evidence_matrix"
    sheet.append(MATRIX_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column, "") for column in MATRIX_COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(MATRIX_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(column) + 2, 12), 28)
    workbook.save(path)


def validate_rows(rows: list[dict[str, str]]) -> dict[str, object]:
    errors: list[dict[str, object]] = []
    warnings: list[dict[str, object]] = []
    required = ["evidence_id", "source_title", "provenance_status", "locator", "parameter_name", "confidence", "parameter_status", "review_decision"]

    for row_number, row in enumerate(rows, start=2):
        for column in required:
            if not row[column]:
                errors.append({"row": row_number, "column": column, "message": "required value is empty"})
        provenance = row["provenance_status"]
        status = row["parameter_status"]
        confidence = row["confidence"]
        decision = row["review_decision"]
        if provenance and provenance not in PROVENANCE:
            errors.append({"row": row_number, "column": "provenance_status", "message": f"unsupported value: {provenance}"})
        if status and status not in PARAMETER_STATUS:
            errors.append({"row": row_number, "column": "parameter_status", "message": f"unsupported value: {status}"})
        if confidence and confidence not in CONFIDENCE:
            errors.append({"row": row_number, "column": "confidence", "message": f"unsupported value: {confidence}"})
        if decision and decision not in REVIEW_DECISIONS:
            errors.append({"row": row_number, "column": "review_decision", "message": f"unsupported value: {decision}"})

        verified = provenance in {"FULL-TEXT VERIFIED", "SUPPLEMENT VERIFIED"}
        if verified and not _truthy(row["full_text_checked"]):
            errors.append({"row": row_number, "column": "full_text_checked", "message": "verified provenance requires full_text_checked=true"})
        if status == "EVIDENCE-SUPPORTED" and not verified:
            errors.append({"row": row_number, "column": "parameter_status", "message": "EVIDENCE-SUPPORTED requires full-text or supplement verified provenance"})
        if status == "EVIDENCE-SUPPORTED" and not row["locator"]:
            errors.append({"row": row_number, "column": "locator", "message": "parameter evidence needs a source locator"})
        if provenance in {"ABSTRACT ONLY", "SECONDARY CITATION", "UNVERIFIED"} and status not in {"", "UNRESOLVED", "PILOT-REQUIRED", "NOT-APPLICABLE"}:
            warnings.append({"row": row_number, "column": "provenance_status", "message": "discovery-only provenance should not drive an executable parameter"})
        if confidence == "INSUFFICIENT" and status not in {"PILOT-REQUIRED", "UNRESOLVED", "NOT-APPLICABLE"}:
            warnings.append({"row": row_number, "column": "confidence", "message": "INSUFFICIENT confidence should normally be pilot-required or unresolved"})

    return {
        "schema": "preclinical-experimental-design/evidence-matrix/v1",
        "rows": len(rows),
        "errors": errors,
        "warnings": warnings,
        "valid": not errors,
    }


def command_init(args: argparse.Namespace) -> int:
    rows: list[dict[str, str]] = []
    write_csv(Path(args.csv), rows)
    if args.xlsx:
        write_xlsx(Path(args.xlsx), rows)
    return 0


def command_validate(args: argparse.Namespace) -> int:
    try:
        rows = read_csv(Path(args.input))
        report = validate_rows(rows)
    except (OSError, ValueError) as exc:
        report = {"valid": False, "rows": 0, "errors": [{"message": str(exc)}], "warnings": []}
    output = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        Path(args.report).write_text(output + "\n", encoding="utf-8")
    print(output)
    return 0 if report["valid"] else 1


def command_export_xlsx(args: argparse.Namespace) -> int:
    rows = read_csv(Path(args.input))
    write_xlsx(Path(args.output), rows)
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    init = subparsers.add_parser("init", help="create blank CSV and optional XLSX templates")
    init.add_argument("--csv", required=True)
    init.add_argument("--xlsx")
    init.set_defaults(func=command_init)
    validate = subparsers.add_parser("validate", help="validate a CSV against the evidence schema")
    validate.add_argument("--input", required=True)
    validate.add_argument("--report")
    validate.set_defaults(func=command_validate)
    export = subparsers.add_parser("export-xlsx", help="export a canonical CSV to XLSX")
    export.add_argument("--input", required=True)
    export.add_argument("--output", required=True)
    export.set_defaults(func=command_export_xlsx)
    return parser


if __name__ == "__main__":
    try:
        parsed = build_parser().parse_args()
        raise SystemExit(parsed.func(parsed))
    except RuntimeError as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(2)
